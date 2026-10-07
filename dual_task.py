"""
Dual Task Application
- Mouse tracking task with random bubble movement
- Secondary task: 4-digit code entry
"""

import tkinter as tk
from tkinter import ttk, font
import random
import math
import time
import os
import datetime
import sys
import ctypes
import openpyxl
from openpyxl.styles import Font as XLFont


def _resource_path(relative):
    """Return absolute path to a resource, works for dev and PyInstaller bundle."""
    base = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, relative)


def _load_passcode():
    """Read the parameter-unlock passcode from passcode.txt (kept out of git).
    Returns None if the file is missing or empty."""
    try:
        with open(_resource_path("passcode.txt"), encoding="utf-8") as f:
            return f.read().strip() or None
    except OSError:
        return None

# Tell Windows this process handles DPI itself so that winfo_screenwidth/height
# and winfo_pointerx/y return true physical pixels (e.g. 1920x1080) rather than
# scaled logical pixels (e.g. 1280x720 at 150% scaling).
# This must happen before any Tk window is created.
# We use SetProcessDpiAwareness(1) — "system aware" — because (2) per-monitor
# can cause coordinate issues with overrideredirect.  Value 1 is sufficient to
# get correct screen metrics on a single-monitor setup.
try:
    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass


# ─────────────────────────────────────────────────────────────
#  App-wide constants / defaults
# ─────────────────────────────────────────────────────────────
BUBBLE_RADIUS = 25          # pixels (diameter = 50)
TAIL_LENGTH = 12            # number of tail circles
TAIL_FADE_STEP = 0.07       # opacity step per tail segment
TICK_MS = 20                # ~50 fps
CODE_DIGITS = 4
COUNTDOWN_FROM = 5
CODE_WINDOW = 5.0       # seconds each code is displayed before timing out

DEFAULTS = dict(
    duration=110,
    num_alarms=20,
    alarm_variance=5,
    max_speed=6.0,
    min_speed=2.0,
    mouse_gain=1.0,
    x_area=1000,
    y_area=600,
)


# ─────────────────────────────────────────────────────────────
#  Colour helpers
# ─────────────────────────────────────────────────────────────
def _hex(r, g, b):
    return f"#{r:02x}{g:02x}{b:02x}"

BG       = "#000000"
FG       = "#e0e0e0"
ACCENT   = "#00cc66"
ACCENT2  = "#0099ff"
BTN_BG   = "#1a1a2e"
BTN_HOV  = "#16213e"
ERR      = "#ff4444"
PANEL_BG = "#0d0d1a"
ENTRY_BG = "#111122"
ENTRY_FG = "#ffffff"
SEP      = "#333355"


# ─────────────────────────────────────────────────────────────
#  Reusable styled widgets
# ─────────────────────────────────────────────────────────────
def styled_label(parent, text, size=12, bold=False, color=FG, **kwargs):
    weight = "bold" if bold else "normal"
    return tk.Label(parent, text=text, bg=PANEL_BG, fg=color,
                    font=("Segoe UI", size, weight), **kwargs)


def styled_entry(parent, textvariable=None, width=12):
    e = tk.Entry(parent, textvariable=textvariable, width=width,
                 bg=ENTRY_BG, fg=ENTRY_FG, insertbackground=ENTRY_FG,
                 relief="flat", font=("Segoe UI", 12),
                 highlightbackground=SEP, highlightthickness=1,
                 highlightcolor=ACCENT)
    return e


def styled_button(parent, text, command, width=18, size=12, color=ACCENT):
    b = tk.Button(parent, text=text, command=command,
                  bg=BTN_BG, fg=color, activebackground=BTN_HOV,
                  activeforeground=color, relief="flat",
                  font=("Segoe UI", size, "bold"), width=width,
                  cursor="hand2", pady=8,
                  highlightthickness=1, highlightbackground=color)

    def on_enter(e):
        b.configure(bg=BTN_HOV)

    def on_leave(e):
        b.configure(bg=BTN_BG)

    b.bind("<Enter>", on_enter)
    b.bind("<Leave>", on_leave)
    return b


def separator(parent):
    return tk.Frame(parent, bg=SEP, height=1)


def id_header(subject_id, session_id):
    """'Subject: X   |   Session: Y' — the session part is omitted when blank."""
    text = f"Subject: {subject_id}"
    if session_id:
        text += f"   |   Session: {session_id}"
    return text


# ─────────────────────────────────────────────────────────────
#  Main application controller
# ─────────────────────────────────────────────────────────────
class DualTaskApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Dual Task")
        self.configure(bg=PANEL_BG)
        self.resizable(False, False)
        try:
            self.iconbitmap(_resource_path(os.path.join('images', 'neuro_logo.ico')))
        except Exception:
            pass

        # Scale factor relative to 1920×1080 baseline.
        # Ensures windows and canvas elements look proportionally the same
        # on high-DPI displays (e.g. Surface Pro 2880×1920 → scale ≈ 1.78).
        self.scale = self.winfo_screenheight() / 1080

        # Shared session state
        self.subject_id = tk.StringVar()
        self.session_id = tk.StringVar()
        self.params = {k: tk.DoubleVar(value=v) for k, v in DEFAULTS.items()}
        self.params["duration"] = tk.IntVar(value=DEFAULTS["duration"])
        self.params["num_alarms"] = tk.IntVar(value=DEFAULTS["num_alarms"])
        self.params["alarm_variance"] = tk.IntVar(value=DEFAULTS["alarm_variance"])
        self.buzzer_on = tk.BooleanVar(value=False)
        self.task_mode = tk.StringVar(value="tracking_only")
        self.trial_counter = 0
        _default_save = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
        self.save_dir = tk.StringVar(value=_default_save)

        # Subjective scales — collected once per subject login
        self.sss_score = tk.IntVar(value=0)   # Stanford Sleepiness Scale 1-7
        self.sfs_score = tk.IntVar(value=0)   # Subjective Fatigue Scale 1-7
        self.mouse_hand = tk.StringVar(value="")  # "Right" or "Left" — hand used for mouse
        self._scales_collected = False

        self._frame = None
        self.show_login()

    # ── frame switching ──────────────────────────────────────
    def _set_frame(self, frame_cls, **kwargs):
        if self._frame:
            self._frame.destroy()
        self._frame = frame_cls(self, **kwargs)
        self._frame.pack(fill="both", expand=True)

    def show_login(self):
        self._set_frame(LoginScreen)

    def show_scales(self):
        self._set_frame(ScalesScreen)

    def show_params(self):
        self._set_frame(ParamsScreen)

    def show_task(self):
        self._set_frame(TaskScreen)

    def show_notes(self, results):
        self._set_frame(NotesScreen, results=results)

    def show_results(self, results):
        self._set_frame(ResultsScreen, results=results)


# ─────────────────────────────────────────────────────────────
#  Screen 1 – Login
# ─────────────────────────────────────────────────────────────
class LoginScreen(tk.Frame):
    def __init__(self, master):
        super().__init__(master, bg=PANEL_BG)
        self._build()

    def _build(self):
        self.columnconfigure(0, weight=1)

        # Title
        tk.Label(self, text="DUAL TASK", bg=PANEL_BG, fg=ACCENT,
                 font=("Segoe UI", 32, "bold")).grid(row=0, column=0, pady=(50, 4))
        tk.Label(self, text="Tracking & Secondary Coding Task",
                 bg=PANEL_BG, fg=FG, font=("Segoe UI", 13)).grid(row=1, column=0, pady=(0, 40))

        separator(self).grid(row=2, column=0, sticky="ew", padx=80, pady=(0, 30))

        # Form box
        form = tk.Frame(self, bg=PANEL_BG)
        form.grid(row=3, column=0)

        styled_label(form, "Subject ID", size=12, bold=True, color=FG).grid(
            row=0, column=0, sticky="w", padx=8, pady=(0, 4))
        self._subj_entry = styled_entry(form, textvariable=self.master.subject_id, width=24)
        self._subj_entry.grid(row=1, column=0, padx=8, pady=(0, 18), ipady=6)

        styled_label(form, "Session ID (optional)", size=12, bold=True, color=FG).grid(
            row=2, column=0, sticky="w", padx=8, pady=(0, 4))
        self._sess_entry = styled_entry(form, textvariable=self.master.session_id, width=24)
        self._sess_entry.grid(row=3, column=0, padx=8, pady=(0, 8), ipady=6)

        self._err_label = tk.Label(form, text="", bg=PANEL_BG, fg=ERR,
                                   font=("Segoe UI", 10))
        self._err_label.grid(row=4, column=0)

        separator(self).grid(row=4, column=0, sticky="ew", padx=80, pady=30)

        btn = styled_button(self, "CONTINUE  →", self._on_continue, width=22, size=13)
        btn.grid(row=5, column=0, pady=(0, 30))

        tk.Label(self,
                 text="NASA Neuroscience Lab  |  Built by Austin Bollinger  2026",
                 bg=PANEL_BG, fg="#444466", font=("Segoe UI", 9)).grid(
            row=6, column=0, pady=(0, 20))

        # Bind Enter key
        self._subj_entry.bind("<Return>", lambda e: self._sess_entry.focus())
        self._sess_entry.bind("<Return>", lambda e: self._on_continue())

        self._subj_entry.focus()

        # Window sizing
        s = self.master.scale
        self.master.geometry(f"{int(520*s)}x{int(590*s)}")
        self.master.update_idletasks()
        self._centre_window()

    def _centre_window(self):
        s = self.master.scale
        w, h = int(520 * s), int(590 * s)
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()
        x = (sw - w) // 2
        y = (sh - h) // 2
        self.master.geometry(f"{w}x{h}+{x}+{y}")

    def _on_continue(self):
        subj = self.master.subject_id.get().strip()
        if not subj:
            self._err_label.configure(text="Subject ID is required.")
            self._subj_entry.focus()
            return
        self.master.session_id.set(self.master.session_id.get().strip())
        # Reset scales, task mode, and trial counter for new subject
        self.master._scales_collected = False
        self.master.sss_score.set(0)
        self.master.sfs_score.set(0)
        self.master.mouse_hand.set("")
        self.master.task_mode.set("tracking_only")
        self.master.trial_counter = 0
        self.master.show_scales()


# ─────────────────────────────────────────────────────────────
#  Screen 1b – Subjective Scales
# ─────────────────────────────────────────────────────────────
SSS_ITEMS = [
    (1, "Feeling active and vital; alert; wide awake"),
    (2, "Functioning at a high level, but not at peak; able to concentrate"),
    (3, "Relaxed; awake; not at full alertness; responsive"),
    (4, "A little foggy; not at peak; let down"),
    (5, "Fogginess; beginning to lose interest in remaining awake; slowed down"),
    (6, "Sleepiness; prefer to be lying down; fighting sleep; woozy"),
    (7, "Almost in reverie; sleep onset soon; lost struggle to remain awake"),
]
SFS_ITEMS = [
    (1, "Fully alert; Wide awake"),
    (2, "Very lively; Responsive; But not at peak"),
    (3, "Okay; Somewhat fresh"),
    (4, "A little tired; less than peak"),
    (5, "Moderately tired; let down"),
    (6, "Extremely tired; Very difficult to concentrate"),
    (7, "Completely exhausted; Unable to function effectively; Ready to drop"),
]


class ScalesScreen(tk.Frame):
    def __init__(self, master):
        super().__init__(master, bg=PANEL_BG)
        self._build()

    def _build(self):
        self.columnconfigure(0, weight=1)

        tk.Label(self, text="SUBJECTIVE RATINGS", bg=PANEL_BG, fg=ACCENT,
                 font=("Segoe UI", 24, "bold")).grid(row=0, column=0, pady=(28, 2))
        tk.Label(self,
                 text=id_header(self.master.subject_id.get(),
                                self.master.session_id.get()),
                 bg=PANEL_BG, fg=FG, font=("Segoe UI", 13)).grid(row=1, column=0, pady=(0, 12))
        separator(self).grid(row=2, column=0, sticky="ew", padx=60, pady=(0, 16))

        body = tk.Frame(self, bg=PANEL_BG)
        body.grid(row=3, column=0, padx=60, sticky="ew")
        body.columnconfigure(0, weight=1)
        body.columnconfigure(1, weight=1)

        self._build_scale_block(body, col=0,
                                title="Stanford Sleepiness Scale",
                                citation="(Hoddes et al., Psychophysiol, 1973)",
                                items=SSS_ITEMS,
                                var=self.master.sss_score)

        tk.Frame(body, bg=SEP, width=1).grid(row=0, column=1, rowspan=9,
                                              sticky="ns", padx=20)

        self._build_scale_block(body, col=2,
                                title="Subjective Fatigue Scale",
                                citation="(Samm-Perelli Crew Status Check)",
                                items=SFS_ITEMS,
                                var=self.master.sfs_score)

        separator(self).grid(row=4, column=0, sticky="ew", padx=60, pady=16)

        hand_frame = tk.Frame(self, bg=PANEL_BG)
        hand_frame.grid(row=5, column=0, pady=(0, 4))
        tk.Label(hand_frame,
                 text="Which hand do you normally use to operate a computer mouse?",
                 bg=PANEL_BG, fg=ACCENT2, font=("Segoe UI", 13, "bold")).pack(pady=(0, 8))
        rb_row = tk.Frame(hand_frame, bg=PANEL_BG)
        rb_row.pack()
        for label in ("Left", "Right"):
            # indicatoron=False renders this as a toggle button rather than a
            # native radio circle — Windows theming otherwise draws the
            # unselected circle in a way that can look "filled" on a dark
            # background, making both options appear selected before any
            # choice is made.
            tk.Radiobutton(rb_row, text=label,
                            variable=self.master.mouse_hand, value=label,
                            indicatoron=False, width=10,
                            bg=BTN_BG, fg=FG, selectcolor="#1f4d3a",
                            activebackground=BTN_HOV, activeforeground=ACCENT,
                            font=("Segoe UI", 12, "bold"), relief="flat", bd=0,
                            highlightthickness=1, highlightbackground=SEP,
                            cursor="hand2", pady=8).pack(side="left", padx=10)

        separator(self).grid(row=6, column=0, sticky="ew", padx=60, pady=16)

        self._err_label = tk.Label(self, text="", bg=PANEL_BG, fg=ERR,
                                   font=("Segoe UI", 12))
        self._err_label.grid(row=7, column=0)

        btn_frame = tk.Frame(self, bg=PANEL_BG)
        btn_frame.grid(row=8, column=0, pady=(4, 28))
        styled_button(btn_frame, "← Back", self.master.show_login,
                      width=12, size=11, color="#888899").pack(side="left", padx=10)
        styled_button(btn_frame, "CONTINUE  →", self._on_continue,
                      width=18, size=13).pack(side="left", padx=10)

        self._centre_window()

    def _build_scale_block(self, parent, col, title, citation, items, var):
        tk.Label(parent, text=title, bg=PANEL_BG, fg=ACCENT2,
                 font=("Segoe UI", 15, "bold")).grid(
            row=0, column=col, sticky="w", pady=(0, 2))
        tk.Label(parent, text=citation, bg=PANEL_BG, fg="#666680",
                 font=("Segoe UI", 11, "italic")).grid(
            row=1, column=col, sticky="w", pady=(0, 8))
        for i, (val, text) in enumerate(items):
            rb = tk.Radiobutton(parent, text=f"{val}  {text}",
                                variable=var, value=val,
                                bg=PANEL_BG, fg=FG, selectcolor=ENTRY_BG,
                                activebackground=PANEL_BG, activeforeground=ACCENT,
                                font=("Segoe UI", 12), anchor="w",
                                highlightthickness=0, cursor="hand2")
            rb.grid(row=i + 2, column=col, sticky="w", pady=4)

    def _centre_window(self):
        s = self.master.scale
        # Let tkinter compute the natural size of all children, then size the
        # window to fit — this adapts to font rendering differences across DPI.
        self.update_idletasks()
        w = max(self.winfo_reqwidth() + 32, int(940 * s))
        h = max(self.winfo_reqheight() + 32, int(740 * s))
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()
        self.master.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def _on_continue(self):
        if self.master.sss_score.get() == 0:
            self._err_label.configure(text="Please select a Sleepiness rating.")
            return
        if self.master.sfs_score.get() == 0:
            self._err_label.configure(text="Please select a Fatigue rating.")
            return
        if not self.master.mouse_hand.get():
            self._err_label.configure(text="Please select your mouse hand.")
            return
        self.master._scales_collected = True
        self.master.show_params()


# ─────────────────────────────────────────────────────────────
#  Screen 2 – Task Parameters
# ─────────────────────────────────────────────────────────────
class ParamsScreen(tk.Frame):
    _PASSCODE = _load_passcode()

    def __init__(self, master):
        super().__init__(master, bg=PANEL_BG)
        self._unlocked = False
        self._build()

    def _build(self):
        self.columnconfigure(0, weight=1)
        p = self.master.params

        tk.Label(self, text="TASK PARAMETERS", bg=PANEL_BG, fg=ACCENT,
                 font=("Segoe UI", 24, "bold")).grid(row=0, column=0, pady=(24, 4))
        tk.Label(self,
                 text=id_header(self.master.subject_id.get(),
                                self.master.session_id.get()),
                 bg=PANEL_BG, fg=FG, font=("Segoe UI", 11)).grid(row=1, column=0, pady=(0, 14))

        separator(self).grid(row=2, column=0, sticky="ew", padx=60, pady=(0, 14))

        grid = tk.Frame(self, bg=PANEL_BG)
        grid.grid(row=3, column=0, padx=60)

        fields = [
            ("Duration (s)",         "duration",       "int"),
            ("Number of Codes",      "num_alarms",      "int"),
            ("Code Interval ± (s)",  "alarm_variance",  "int"),
            ("Max Speed (px/tick)",  "max_speed",       "float"),
            ("Min Speed (px/tick)",  "min_speed",       "float"),
            ("Mouse Gain",           "mouse_gain",      "float"),
        ]

        self._entries = {}
        for i, (label, key, kind) in enumerate(fields):
            styled_label(grid, label, size=11, color=FG).grid(
                row=i, column=0, sticky="w", padx=(0, 20), pady=6)
            var = p[key]
            e = styled_entry(grid, textvariable=var, width=10)
            e.grid(row=i, column=1, sticky="w", pady=4, ipady=4)
            e.configure(state="disabled")
            self._entries[key] = (e, kind)
            # Buzzer checkbox sits on the same row as "Number of Codes"
            if key == "num_alarms":
                self._buzzer_cb = tk.Checkbutton(grid, text="Play buzzer on code",
                               variable=self.master.buzzer_on,
                               bg=PANEL_BG, fg=FG, selectcolor=ENTRY_BG,
                               activebackground=PANEL_BG, activeforeground=FG,
                               font=("Segoe UI", 10), cursor="hand2", state="disabled"
                               )
                self._buzzer_cb.grid(row=i, column=2, sticky="w", padx=(20, 0))

        # Task Mode selector
        mode_frame = tk.Frame(self, bg=PANEL_BG)
        mode_frame.grid(row=4, column=0, padx=60, sticky="ew", pady=(6, 2))
        styled_label(mode_frame, "Task Mode", size=11, bold=True, color=FG).grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 6))
        self._mode_rbs = []
        for col, (val, label) in enumerate([
            ("tracking_only", "Tracking Only"),
            ("codes_only",    "Codes Only"),
            ("dual_task",     "Dual Task  (Tracking + Codes)"),
        ]):
            rb = tk.Radiobutton(mode_frame, text=label,
                                variable=self.master.task_mode, value=val,
                                bg=PANEL_BG, fg=FG, selectcolor=ENTRY_BG,
                                activebackground=PANEL_BG, activeforeground=ACCENT,
                                font=("Segoe UI", 11), cursor="hand2", state="disabled")
            rb.grid(row=1, column=col, sticky="w", padx=(0, 24))
            self._mode_rbs.append(rb)

        separator(self).grid(row=5, column=0, sticky="ew", padx=60, pady=14)

        # Save path row
        path_frame = tk.Frame(self, bg=PANEL_BG)
        path_frame.grid(row=6, column=0, padx=60, sticky="ew", pady=(0, 6))
        path_frame.columnconfigure(1, weight=1)

        styled_label(path_frame, "Save Folder", size=11, color=FG).grid(
            row=0, column=0, sticky="w", padx=(0, 12))
        path_entry = tk.Entry(path_frame, textvariable=self.master.save_dir,
                              bg=ENTRY_BG, fg=ENTRY_FG, insertbackground=ENTRY_FG,
                              relief="flat", font=("Segoe UI", 10),
                              highlightbackground=SEP, highlightthickness=1,
                              highlightcolor=ACCENT)
        path_entry.grid(row=0, column=1, sticky="ew", ipady=4)
        browse_btn = tk.Button(path_frame, text="Browse", command=self._browse_folder,
                               bg=BTN_BG, fg=ACCENT2, activebackground=BTN_HOV,
                               activeforeground=ACCENT2, relief="flat",
                               font=("Segoe UI", 10), cursor="hand2",
                               highlightthickness=1, highlightbackground=ACCENT2,
                               padx=10, pady=4)
        browse_btn.grid(row=0, column=2, padx=(8, 0))

        # Task instruction box — updates with mode selection and the subject's
        # selected mouse hand (collected on the Subjective Ratings screen).
        def _mode_instructions(mouse_hand):
            # Fall back to generic phrasing if a hand was somehow never recorded.
            hand = mouse_hand or "the hand you normally use to operate a computer mouse"
            other_hand = {"Right": "Left", "Left": "Right"}.get(mouse_hand)
            other = (f"your {other_hand} hand" if other_hand
                     else "the hand opposite the one you normally use to operate a computer mouse")
            mouse_phrase = f"your {hand} hand" if mouse_hand else hand

            return {
                "tracking_only": (
                    "Trial 1 — TRACKING ONLY",
                    [
                        ("Tracking Task:",
                         "A green circle (bubble) will move around the screen. Use the mouse to keep "
                         "the white crosshair cursor inside the bubble for the entire task. Try to "
                         "remain inside the bubble as much as possible and keep the crosshair as close "
                         f"to the center of the bubble as you can. Please use {mouse_phrase} "
                         "to control the mouse."),
                    ]
                ),
                "codes_only": (
                    "Trial 2 — CODES ONLY",
                    [
                        ("Code Entry Task:",
                         "A 4-digit code will appear at the top of the screen (for example, 4829). "
                         f"Use {other} "
                         "to type the four digits as quickly and accurately as possible. Each code "
                         "remains on the screen for 5 seconds before disappearing, regardless of "
                         "whether you have finished entering it. A new code will appear every 5 seconds.\n\n"
                         "A code is scored as correct only if all four digits are entered correctly "
                         "before it disappears. If any digit is incorrect, omitted, or the response "
                         "is not completed within 5 seconds, that code will be scored as incorrect."),
                    ]
                ),
                "dual_task": (
                    "Trial 3 — DUAL TASK",
                    [
                        ("Tracking Task:",
                         "Keep the white crosshair inside the moving bubble using the mouse. Use "
                         f"{mouse_phrase} — the same hand you used during the Tracking Only task."),
                        ("Code Entry Task:",
                         "When a 4-digit code appears at the top of the screen, enter it using "
                         f"{other} — the same hand you used during the Codes Only task — while "
                         "continuing to perform the tracking task. The code will disappear after 5 seconds."),
                        ("Task Priority:",
                         "Give equal priority to both tasks. Try to maintain good tracking performance "
                         "while also entering the codes accurately. Do not intentionally neglect one "
                         "task in favor of the other. Both tasks are equally important. Perform each "
                         "task as accurately as possible while responding as quickly as you can."),
                    ]
                ),
            }

        hint_outer = tk.Frame(self, bg="#1a1a2e",
                              highlightbackground=ACCENT2, highlightthickness=1)
        hint_outer.grid(row=7, column=0, padx=60, pady=(8, 10), sticky="ew")
        hint_outer.columnconfigure(0, weight=1)

        self._hint_title = tk.Label(hint_outer, text="", bg="#1a1a2e", fg=ACCENT2,
                                    font=("Segoe UI", 11, "bold"), anchor="w")
        self._hint_title.grid(row=0, column=0, sticky="ew", padx=14, pady=(10, 4))

        self._hint_text = tk.Text(hint_outer, bg="#1a1a2e", fg=FG,
                                  font=("Segoe UI", 10), relief="flat",
                                  wrap="word", height=10, state="disabled",
                                  cursor="arrow", highlightthickness=0, bd=0,
                                  spacing1=1, spacing2=2, spacing3=6)
        self._hint_text.tag_configure("header", foreground=ACCENT2,
                                      font=("Segoe UI", 10, "bold"))
        self._hint_text.tag_configure("body", foreground=FG,
                                      font=("Segoe UI", 10))
        self._hint_text.grid(row=1, column=0, sticky="ew", padx=14, pady=(0, 10))

        def _update_hint(*_):
            mode = self.master.task_mode.get()
            instructions = _mode_instructions(self.master.mouse_hand.get())
            title, sections = instructions.get(mode, ("", []))
            self._hint_title.configure(text=title)
            self._hint_text.configure(state="normal")
            self._hint_text.delete("1.0", "end")
            for i, (header, body) in enumerate(sections):
                if i > 0:
                    self._hint_text.insert("end", "\n")
                self._hint_text.insert("end", header + "\n", "header")
                self._hint_text.insert("end", body, "body")
            self._hint_text.configure(state="disabled")

        _trace_id = self.master.task_mode.trace_add("write", _update_hint)

        def _remove_trace(_=None):
            try:
                self.master.task_mode.trace_remove("write", _trace_id)
            except Exception:
                pass

        self.bind("<Destroy>", _remove_trace)
        _update_hint()

        # Buttons
        btn_frame = tk.Frame(self, bg=PANEL_BG)
        btn_frame.grid(row=8, column=0, pady=(0, 28))

        styled_button(btn_frame, "← Back", self.master.show_login,
                      width=12, size=11, color="#888899").pack(side="left", padx=10)
        self._unlock_btn = styled_button(btn_frame, "🔒 Unlock Params",
                      self._prompt_unlock, width=18, size=11, color=ACCENT2)
        self._unlock_btn.pack(side="left", padx=10)
        styled_button(btn_frame, "START TASK  →", self._on_start,
                      width=18, size=13).pack(side="left", padx=10)

        s = self.master.scale
        self.master.geometry(f"{int(720*s)}x{int(900*s)}")
        self._centre_window()

    def _prompt_unlock(self):
        dlg = tk.Toplevel(self.master)
        dlg.title("Unlock Parameters")
        dlg.configure(bg=PANEL_BG)
        dlg.resizable(False, False)
        dlg.grab_set()

        tk.Label(dlg, text="Enter passcode:", bg=PANEL_BG, fg=FG,
                 font=("Segoe UI", 12)).pack(padx=30, pady=(24, 6))
        pw_var = tk.StringVar()
        pw_entry = tk.Entry(dlg, textvariable=pw_var, show="*",
                            bg=ENTRY_BG, fg=ENTRY_FG, insertbackground=ENTRY_FG,
                            relief="flat", font=("Segoe UI", 12), width=20,
                            highlightbackground=SEP, highlightthickness=1,
                            highlightcolor=ACCENT)
        pw_entry.pack(padx=30, ipady=6)
        pw_entry.focus()

        err_lbl = tk.Label(dlg, text="", bg=PANEL_BG, fg=ERR,
                           font=("Segoe UI", 10))
        err_lbl.pack(pady=(4, 0))

        def _check(event=None):
            if self._PASSCODE is None:
                err_lbl.configure(text="passcode.txt not found — cannot unlock.")
                pw_var.set("")
            elif pw_var.get() == self._PASSCODE:
                self._unlocked = True
                for e, _ in self._entries.values():
                    e.configure(state="normal")
                self._buzzer_cb.configure(state="normal")
                for rb in self._mode_rbs:
                    rb.configure(state="normal")
                self._unlock_btn.configure(text="🔓 Params Unlocked",
                                           state="disabled", fg="#888899")
                dlg.destroy()
            else:
                err_lbl.configure(text="Incorrect passcode.")
                pw_var.set("")
                pw_entry.focus()

        pw_entry.bind("<Return>", _check)
        tk.Button(dlg, text="Unlock", command=_check,
                  bg=BTN_BG, fg=ACCENT, activebackground=BTN_HOV,
                  activeforeground=ACCENT, relief="flat",
                  font=("Segoe UI", 11, "bold"), cursor="hand2",
                  highlightthickness=1, highlightbackground=ACCENT,
                  padx=16, pady=6).pack(pady=(12, 24))

        # Centre the dialog
        dlg.update_idletasks()
        dw, dh = dlg.winfo_width(), dlg.winfo_height()
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()
        dlg.geometry(f"+{(sw-dw)//2}+{(sh-dh)//2}")

    def _browse_folder(self):
        from tkinter import filedialog
        current = self.master.save_dir.get()
        chosen = filedialog.askdirectory(
            title="Select Save Folder",
            initialdir=current if os.path.isdir(current) else os.path.expanduser("~"))
        if chosen:
            self.master.save_dir.set(chosen)

    def _centre_window(self):
        s = self.master.scale
        w, h = int(720 * s), int(900 * s)
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()
        self.master.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def _on_start(self):
        # Validate
        p = self.master.params
        try:
            dur = int(p["duration"].get())
            n = int(p["num_alarms"].get())
            vr = int(p["alarm_variance"].get())
            mx = float(p["max_speed"].get())
            mn = float(p["min_speed"].get())
            gain = float(p["mouse_gain"].get())
        except (ValueError, tk.TclError):
            return

        if dur <= 0 or n < 0 or mx <= 0 or mn <= 0:
            return
        if mn > mx:
            p["min_speed"].set(mx)

        self.master.show_task()


# ─────────────────────────────────────────────────────────────
#  Bubble physics
# ─────────────────────────────────────────────────────────────
class Bubble:
    """Random-walk bubble with smooth speed variation."""

    def __init__(self, cx, cy, min_speed, max_speed):
        self.x = float(cx)
        self.y = float(cy)
        self.min_speed = min_speed
        self.max_speed = max_speed
        self.speed = (min_speed + max_speed) / 2
        self.angle = random.uniform(0, 2 * math.pi)
        self._speed_target = self.speed
        self._angle_change_rate = 0.05
        # tail positions list[(x, y)]
        self.tail = []

    def update(self, area_w, area_h, top_bound=0):
        # Randomly drift target speed
        if random.random() < 0.03:
            self._speed_target = random.uniform(self.min_speed, self.max_speed)
        self.speed += (self._speed_target - self.speed) * 0.04

        # Slowly drift angle
        self.angle += random.gauss(0, self._angle_change_rate)

        dx = math.cos(self.angle) * self.speed
        dy = math.sin(self.angle) * self.speed

        nx = self.x + dx
        ny = self.y + dy

        # Bounce off walls with margin = radius
        r = BUBBLE_RADIUS
        if nx - r < 0:
            nx = r
            self.angle = math.pi - self.angle
        elif nx + r > area_w:
            nx = area_w - r
            self.angle = math.pi - self.angle

        if ny - r < top_bound:
            ny = top_bound + r
            self.angle = -self.angle
        elif ny + r > area_h:
            ny = area_h - r
            self.angle = -self.angle

        # Save tail
        self.tail.append((self.x, self.y))
        if len(self.tail) > TAIL_LENGTH:
            self.tail.pop(0)

        self.x = nx
        self.y = ny


# ─────────────────────────────────────────────────────────────
#  Screen 3 – Task
# ─────────────────────────────────────────────────────────────
class TaskScreen(tk.Frame):
    def __init__(self, master):
        super().__init__(master, bg=BG)
        self._p = master.params
        self._running = False
        self._counting_down = False
        self._task_started = False
        self._start_time = None

        # Data
        self._tracking_data = []   # (elapsed, cursor_x, cursor_y, bubble_x, bubble_y, error)
        self._code_data = []       # (elapsed, code, response, rt, correct)
        self._cursor_x = 0.0
        self._cursor_y = 0.0

        # Code state
        self._current_code = None
        self._code_start_time = None
        self._code_expire_time = None
        self._next_code_time = None
        self._code_schedule = []
        self._code_index = 0
        self._input_buffer = ""
        self._input_blocked_until = 0.0   # grace period after each new code appears

        self._build()
        self._setup_task()

    # ── Layout ───────────────────────────────────────────────
    def _build(self):
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()

        self._area_w = sw
        self._area_h = sh

        self.master.resizable(True, True)
        self.master.attributes("-fullscreen", True)
        self.master.attributes("-topmost", True)

        self._canvas = tk.Canvas(self, width=sw, height=sh,
                                 bg=BG, highlightthickness=0, cursor="none")
        self._canvas.pack(fill="both", expand=True)

        self._tracking_top = max(80, int(96 * self.master.scale))
        self._status_var = tk.StringVar(value="")
        self.master.bind("<KeyPress>", self._on_key)
        self.master.bind("<Escape>", self._on_escape)

    # ── Setup ─────────────────────────────────────────────────
    def _setup_task(self):

        area_w = self._area_w
        area_h = self._area_h
        p = self._p

        min_s = float(p["min_speed"].get())
        max_s = float(p["max_speed"].get())
        self._bubble = Bubble(area_w // 2, area_h // 2, min_s, max_s)

        self._cursor_x = float(self._bubble.x)
        self._cursor_y = float(self._bubble.y)
        self._gain = float(p["mouse_gain"].get())
        self._duration = int(p["duration"].get())

        # Determine which sub-tasks are active this trial
        self._task_mode = self.master.task_mode.get()
        self._do_tracking = self._task_mode in ("tracking_only", "dual_task")
        self._do_codes = self._task_mode in ("codes_only", "dual_task")

        # Build code schedule — fixed CODE_WINDOW intervals after a 5-second buffer.
        # Codes at t=5, 10, 15, ... Only include codes that start before task end.
        if self._do_codes:
            num_codes = int(p["num_alarms"].get())
            if num_codes > 0:
                sched = [5.0 + i * CODE_WINDOW for i in range(num_codes)
                         if 5.0 + i * CODE_WINDOW < self._duration]
                self._code_schedule = sched
            else:
                self._code_schedule = []
        else:
            self._code_schedule = []

        self._start_countdown()

    # ── Countdown ────────────────────────────────────────────
    def _start_countdown(self):
        self._counting_down = True
        self._countdown_val = COUNTDOWN_FROM
        self._move_pointer_to_tracking_point(self._bubble.x, self._bubble.y)
        self._draw_countdown()

    def _move_pointer_to_tracking_point(self, tracking_x, tracking_y):
        gain = self._gain if self._gain else 1.0
        usable_mid_y = (self._tracking_top + self._area_h) / 2
        ptr_x = self._area_w / 2 + (tracking_x - self._area_w / 2) / gain
        ptr_y = usable_mid_y + (tracking_y - usable_mid_y) / gain
        screen_x = int(round(self.master.winfo_rootx() + ptr_x))
        screen_y = int(round(self.master.winfo_rooty() + ptr_y))
        try:
            ctypes.windll.user32.SetCursorPos(screen_x, screen_y)
        except Exception:
            pass

    def _draw_countdown(self):
        c = self._canvas
        c.delete("all")
        w, h = self._area_w, self._area_h
        s = self.master.scale
        num_font = max(60, int(120 * s))
        sub_font = max(14, int(24 * s))
        # Place number at 40% height, "Get ready..." at 65% height.
        # Using screen fractions keeps them well separated at any resolution.
        num_y = int(h * 0.40)
        sub_y = int(h * 0.65)

        if self._do_tracking:
            br = max(8, int(BUBBLE_RADIUS * s))
            # ── Bubble ──
            bx, by = self._bubble.x, self._bubble.y
            c.create_oval(bx - br, by - br, bx + br, by + br,
                          outline=ACCENT, width=max(1, int(2 * s)), fill="")
            # ── Crosshair ──
            cx, cy = self._cursor_x, self._cursor_y
            arm = max(8, int(14 * s))
            gap = max(3, int(5 * s))
            lw  = max(1, int(2 * s))
            c.create_line(cx - arm, cy, cx - gap, cy, fill=ENTRY_FG, width=lw)
            c.create_line(cx + gap, cy, cx + arm, cy, fill=ENTRY_FG, width=lw)
            c.create_line(cx, cy - arm, cx, cy - gap, fill=ENTRY_FG, width=lw)
            c.create_line(cx, cy + gap, cx, cy + arm, fill=ENTRY_FG, width=lw)

        if self._countdown_val > 0:
            c.create_text(
                w // 2, num_y,
                text=str(self._countdown_val),
                fill=ACCENT, font=("Segoe UI", num_font, "bold"))
            c.create_text(
                w // 2, sub_y,
                text="Get ready...",
                fill=FG, font=("Segoe UI", sub_font))
            self._countdown_val -= 1
            self.after(1000, self._draw_countdown)
        else:
            c.create_text(
                w // 2, num_y,
                text="GO!",
                fill=ACCENT, font=("Segoe UI", num_font, "bold"))
            self.after(800, self._begin_task)

    # ── Task start ───────────────────────────────────────────
    def _begin_task(self):
        self._move_pointer_to_tracking_point(self._bubble.x, self._bubble.y)
        self._counting_down = False
        self._task_started = True
        self._running = True
        self._start_time = time.perf_counter()
        self._code_index = 0
        self._next_code_time = self._code_schedule[0] if self._code_schedule else None
        self._tick()

    # ── Main animation loop ──────────────────────────────────
    def _tick(self):
        if not self._running:
            return

        elapsed = time.perf_counter() - self._start_time

        if elapsed >= self._duration:
            self._end_task()
            return

        p = self._p

        if self._do_tracking:
            gain = float(p["mouse_gain"].get())
            # Direct position mapping: cursor = mouse position on screen.
            # Gain scales sensitivity around the screen centre: 1.0 = 1-to-1,
            # <1.0 = reduced range (easier), >1.0 = amplified (harder).
            ptr_x = self.master.winfo_pointerx() - self.master.winfo_rootx()
            ptr_y = self.master.winfo_pointery() - self.master.winfo_rooty()
            cx = self._area_w / 2 + (ptr_x - self._area_w / 2) * gain
            usable_mid_y = (self._tracking_top + self._area_h) / 2
            cy = usable_mid_y + (ptr_y - usable_mid_y) * gain
            self._cursor_x = max(0.0, min(float(self._area_w), cx))
            self._cursor_y = max(float(self._tracking_top), min(float(self._area_h), cy))

            self._bubble.update(self._area_w, self._area_h, self._tracking_top)
            err = math.hypot(self._cursor_x - self._bubble.x,
                             self._cursor_y - self._bubble.y)
            self._tracking_data.append((
                round(elapsed, 3),
                round(self._cursor_x, 1),
                round(self._cursor_y, 1),
                round(self._bubble.x, 1),
                round(self._bubble.y, 1),
                round(err, 2),
            ))

        if self._do_codes:
            # Expire current code if its 5-second window has passed (marks it MISSED)
            if self._current_code is not None and elapsed >= self._code_expire_time:
                self._miss_code(elapsed)
            # Show next scheduled code
            if (self._next_code_time is not None and
                    elapsed >= self._next_code_time and
                    self._current_code is None):
                self._show_code(elapsed)

        self._draw_frame(elapsed)
        self.after(TICK_MS, self._tick)

    # ── Draw ─────────────────────────────────────────────────
    def _draw_frame(self, elapsed):
        c = self._canvas
        c.delete("all")
        w, h = self._area_w, self._area_h
        s = self.master.scale
        br = max(8, int(BUBBLE_RADIUS * s))  # scaled bubble radius

        if self._do_tracking:
            # ── Bubble tail ──
            tail = self._bubble.tail
            for i, (tx, ty) in enumerate(tail):
                if i % 2 == 1 and 0 < i < len(tail) - 1:
                    continue
                frac = (i + 1) / (len(tail) + 1)
                r_size = max(4, int(br * frac * 0.85))
                grey = int(40 + 160 * frac)
                col = _hex(grey * 2 // 3, grey, grey * 2 // 3)
                c.create_oval(tx - r_size, ty - r_size,
                              tx + r_size, ty + r_size,
                              outline=col, width=1, fill="")

            # ── Bubble ──
            bx, by = self._bubble.x, self._bubble.y
            c.create_oval(bx - br, by - br, bx + br, by + br,
                          outline=ACCENT, width=max(1, int(2 * s)), fill="")

            # ── Cursor (crosshair) ──
            cx, cy = self._cursor_x, self._cursor_y
            arm = max(8, int(14 * s))
            gap = max(3, int(5 * s))
            lw  = max(1, int(2 * s))
            c.create_line(cx - arm, cy, cx - gap, cy, fill=ENTRY_FG, width=lw)
            c.create_line(cx + gap, cy, cx + arm, cy, fill=ENTRY_FG, width=lw)
            c.create_line(cx, cy - arm, cx, cy - gap, fill=ENTRY_FG, width=lw)
            c.create_line(cx, cy + gap, cx, cy + arm, fill=ENTRY_FG, width=lw)

        # ── Progress bar (thin, top edge) ──
        progress = min(elapsed / self._duration, 1.0)
        bar_h = max(2, int(3 * s))
        c.create_rectangle(0, 0, int(w * progress), bar_h,
                            fill=ACCENT2, outline="")

        # ── Elapsed / duration ──
        remaining = max(0, self._duration - elapsed)
        timer_y = max(10, int(14 * s))
        timer_font = max(8, int(11 * s))
        c.create_text(w - int(16 * s), timer_y, text=f"{int(remaining)}s",
                      fill="#555577", font=("Segoe UI", timer_font), anchor="ne")

        # ── Code prompt ──
        if self._current_code is not None:
            code_str = str(self._current_code)
            padded_input = self._input_buffer.ljust(CODE_DIGITS, "_")
            prompt = f"Enter: {code_str}      Your input: {padded_input}"
            prompt_font = max(14, int(28 * s))
            prompt_y = max(20, int(36 * s))
            c.create_text(w // 2, prompt_y,
                          text=prompt,
                          fill=ACCENT, font=("Segoe UI", prompt_font, "bold"),
                          anchor="center")

    # ── Code logic ───────────────────────────────────────────
    def _show_code(self, elapsed):
        self._current_code = "".join([str(random.randint(0, 9))
                                      for _ in range(CODE_DIGITS)])
        self._code_start_time = elapsed
        self._code_expire_time = elapsed + CODE_WINDOW
        self._input_buffer = ""
        self._input_blocked_until = elapsed + 0.3  # drain stale queued keypresses
        if self.master.buzzer_on.get():
            import threading, winsound
            threading.Thread(target=lambda: winsound.Beep(880, 180),
                             daemon=True).start()
        # Advance schedule
        self._code_index += 1
        if self._code_index < len(self._code_schedule):
            self._next_code_time = self._code_schedule[self._code_index]
        else:
            self._next_code_time = None

    def _miss_code(self, elapsed):
        """Record the current code as missed — window expired without a response."""
        if self._current_code is None:
            return
        rt = elapsed - self._code_start_time
        self._code_data.append((
            round(self._code_start_time, 3),
            self._current_code,
            "MISSED",
            round(rt, 3),
            0,
        ))
        self._current_code = None
        self._code_expire_time = None
        self._input_buffer = ""

    def _submit_code(self, elapsed):
        if self._current_code is None:
            return
        rt = elapsed - self._code_start_time
        correct = (self._input_buffer == self._current_code)
        self._code_data.append((
            round(self._code_start_time, 3),
            self._current_code,
            self._input_buffer,
            round(rt, 3),
            int(correct),
        ))
        self._current_code = None
        self._code_expire_time = None
        self._input_buffer = ""

    def _dismiss_code(self, elapsed):
        """Record an incomplete code when the task ends before submission."""
        if self._current_code is None:
            return
        rt = elapsed - self._code_start_time
        self._code_data.append((
            round(self._code_start_time, 3),
            self._current_code,
            self._input_buffer if self._input_buffer else "NO_RESPONSE",
            round(rt, 3),
            0,
        ))
        self._current_code = None
        self._code_expire_time = None
        self._input_buffer = ""

    # ── Input ────────────────────────────────────────────────
    def _on_key(self, event):
        if not self._task_started or not self._running:
            return
        elapsed = time.perf_counter() - self._start_time
        key = event.keysym

        if elapsed < self._input_blocked_until:
            return
        if key.isdigit() and self._current_code is not None:
            if len(self._input_buffer) < CODE_DIGITS:
                self._input_buffer += key
                if len(self._input_buffer) == CODE_DIGITS:
                    self._submit_code(elapsed)
        elif key == "Return" and self._current_code is not None:
            self._submit_code(elapsed)

    def _on_escape(self, event):
        if self._running:
            self._end_task()

    # ── End task ─────────────────────────────────────────────
    def _end_task(self):
        self._running = False
        elapsed = time.perf_counter() - self._start_time

        # Dismiss any pending code
        if self._current_code is not None:
            self._dismiss_code(elapsed)

        # Compute results
        results = self._compute_results()
        # Save immediately so the data is safe even if the app is closed on
        # the notes screen; the notes screen re-saves the same file.
        save_results(self.master, results)

        # Restore window to normal managed mode
        self.master.attributes("-fullscreen", False)
        self.master.attributes("-topmost", False)
        self.master.resizable(False, False)
        self.master.show_notes(results)

    def _compute_results(self):
        dur = int(self._p["duration"].get())
        task_parameters = {k: v.get() for k, v in self._p.items()}
        task_parameters.update({
            "task_mode": self._task_mode,
            "buzzer_on": self.master.buzzer_on.get(),
            "bubble_radius_px": BUBBLE_RADIUS,
            "tail_length": TAIL_LENGTH,
            "tick_ms": TICK_MS,
            "code_digits": CODE_DIGITS,
            "countdown_from": COUNTDOWN_FROM,
        })

        # Tracking
        if self._tracking_data:
            errors = [row[5] for row in self._tracking_data]
            avg_err = sum(errors) / len(errors)
            max_err = max(errors)
            on_target = sum(1 for e in errors if e <= BUBBLE_RADIUS * self.master.scale)
            time_on_target = round(on_target / len(errors) * 100, 1)
        else:
            avg_err = max_err = 0.0
            time_on_target = 0.0

        # Coding
        total_codes = len(self._code_data)
        correct_codes = sum(r[4] for r in self._code_data)
        incorrect_codes = total_codes - correct_codes
        accuracy = (correct_codes / total_codes * 100) if total_codes else 0.0
        rts = [r[3] for r in self._code_data if r[4] == 1]
        avg_rt = sum(rts) / len(rts) if rts else 0.0

        return dict(
            subject_id=self.master.subject_id.get(),
            session_id=self.master.session_id.get(),
            task_mode=self._task_mode,
            sss_scale=self.master.sss_score.get(),
            sfs_scale=self.master.sfs_score.get(),
            mouse_hand=self.master.mouse_hand.get(),
            duration=dur,
            avg_tracking_error=round(avg_err, 2),
            max_tracking_error=round(max_err, 2),
            time_on_target=time_on_target,
            total_codes=total_codes,
            correct_codes=correct_codes,
            incorrect_codes=incorrect_codes,
            code_accuracy=round(accuracy, 1),
            avg_correct_rt=round(avg_rt, 3),
            notes="",
            task_parameters=task_parameters,
            tracking_data=self._tracking_data,
            code_data=self._code_data,
        )


# ─────────────────────────────────────────────────────────────
#  Saving
# ─────────────────────────────────────────────────────────────
def save_results(app, results):
    """Write results to an .xlsx file.  The first call assigns the trial number
    and file name; later calls (e.g. after notes are entered) overwrite it."""
    mode = results.get("task_mode", "dual_task")

    if "base_name" not in results:
        subj = results["subject_id"]
        sess = results["session_id"]
        app.trial_counter += 1
        trial = app.trial_counter
        now = datetime.datetime.now()
        timestamp = now.strftime("%m_%d_%H%M")
        sess_part = f"_session{sess}" if sess else ""
        results["completed_datetime"] = now.strftime("%m/%d/%Y %H:%M")
        results["base_name"] = f"{subj}{sess_part}_{mode}_trial{trial}_{timestamp}"
        results["save_dir"] = app.save_dir.get().strip() or \
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
    base_name = results["base_name"]
    save_dir = results["save_dir"]
    os.makedirs(save_dir, exist_ok=True)

    wb = openpyxl.Workbook()
    bold = XLFont(bold=True)

    # ── Data tab ─────────────────────────────────────────
    ws_data = wb.active
    ws_data.title = "Data"

    _TRK_COLS = ["elapsed_s", "cursor_x", "cursor_y", "bubble_x", "bubble_y", "error_px"]
    _COD_COLS = ["elapsed_s", "code", "response", "rt_s", "correct"]

    if mode == "dual_task":
        # Tracking block: cols A-F  |  gap col G  |  Code block: cols H-L
        trk_start = 1
        cod_start = len(_TRK_COLS) + 2          # col 8 = H
        for c, h in enumerate(_TRK_COLS, trk_start):
            ws_data.cell(row=1, column=c, value=h).font = bold
        for c, h in enumerate(_COD_COLS, cod_start):
            ws_data.cell(row=1, column=c, value=h).font = bold
        for r, row in enumerate(results["tracking_data"], start=2):
            for c, val in enumerate([row[0], row[1], row[2], row[3], row[4], row[5]], trk_start):
                ws_data.cell(row=r, column=c, value=val)
        for r, row in enumerate(results["code_data"], start=2):
            for c, val in enumerate([row[0], row[1], row[2], row[3], row[4]], cod_start):
                ws_data.cell(row=r, column=c, value=val)

    elif mode == "tracking_only":
        for c, h in enumerate(_TRK_COLS, 1):
            ws_data.cell(row=1, column=c, value=h).font = bold
        for r, row in enumerate(results["tracking_data"], start=2):
            for c, val in enumerate([row[0], row[1], row[2], row[3], row[4], row[5]], 1):
                ws_data.cell(row=r, column=c, value=val)

    else:  # codes_only
        for c, h in enumerate(_COD_COLS, 1):
            ws_data.cell(row=1, column=c, value=h).font = bold
        for r, row in enumerate(results["code_data"], start=2):
            for c, val in enumerate([row[0], row[1], row[2], row[3], row[4]], 1):
                ws_data.cell(row=r, column=c, value=val)

    # ── Summary tab ──────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")
    summary_fields = {"completed_datetime": results["completed_datetime"]}
    for k, v in results.items():
        if k not in ("tracking_data", "code_data", "task_parameters",
                     "completed_datetime", "base_name", "save_dir"):
            summary_fields[k] = str(v)
    ws_sum.append(list(summary_fields.keys()))
    for cell in ws_sum[1]:
        cell.font = bold
    ws_sum.append(list(summary_fields.values()))

    # ── Parameters tab ───────────────────────────────────
    ws_params = wb.create_sheet("Parameters")
    params = results["task_parameters"]
    ws_params.append(list(params.keys()))
    for cell in ws_params[1]:
        cell.font = bold
    ws_params.append([str(v) for v in params.values()])

    file_path = os.path.join(save_dir, f"{base_name}.xlsx")
    wb.save(file_path)


# ─────────────────────────────────────────────────────────────
#  Screen 3b – Post-task Notes
# ─────────────────────────────────────────────────────────────
class NotesScreen(tk.Frame):
    def __init__(self, master, results):
        super().__init__(master, bg=PANEL_BG)
        self._results = results
        self._build()

    def _build(self):
        self.columnconfigure(0, weight=1)

        tk.Label(self, text="TASK COMPLETE", bg=PANEL_BG, fg=ACCENT,
                 font=("Segoe UI", 26, "bold")).grid(row=0, column=0, pady=(36, 4))
        tk.Label(self,
                 text=id_header(self._results["subject_id"], self._results["session_id"]),
                 bg=PANEL_BG, fg=FG, font=("Segoe UI", 11)).grid(row=1, column=0, pady=(0, 20))

        separator(self).grid(row=2, column=0, sticky="ew", padx=60, pady=(0, 18))

        tk.Label(self,
                 text="Do you have any notes or remarks about this task? (optional)",
                 bg=PANEL_BG, fg=ACCENT2, font=("Segoe UI", 13, "bold")).grid(
            row=3, column=0, padx=60, pady=(0, 10))

        self._text = tk.Text(self, width=60, height=8, wrap="word",
                             bg=ENTRY_BG, fg=ENTRY_FG, insertbackground=ENTRY_FG,
                             relief="flat", font=("Segoe UI", 12),
                             highlightbackground=SEP, highlightthickness=1,
                             highlightcolor=ACCENT, padx=8, pady=6)
        self._text.grid(row=4, column=0, padx=60, sticky="ew")

        separator(self).grid(row=5, column=0, sticky="ew", padx=60, pady=18)

        styled_button(self, "CONTINUE  →", self._on_continue,
                      width=22, size=13).grid(row=6, column=0, pady=(0, 36))

        self._centre_window()
        self._text.focus_set()

    def _centre_window(self):
        s = self.master.scale
        self.update_idletasks()
        w = max(self.winfo_reqwidth() + 32, int(780 * s))
        h = max(self.winfo_reqheight() + 32, int(500 * s))
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()
        self.master.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def _on_continue(self):
        notes = self._text.get("1.0", "end").strip()
        if notes:
            self._results["notes"] = notes
            save_results(self.master, self._results)
        self.master.show_results(self._results)


# ─────────────────────────────────────────────────────────────
#  Screen 4 – Results
# ─────────────────────────────────────────────────────────────
class ResultsScreen(tk.Frame):
    def __init__(self, master, results):
        super().__init__(master, bg=PANEL_BG)
        self._results = results
        self._build()

    def _build(self):
        r = self._results
        self.columnconfigure(0, weight=1)

        # Title
        tk.Label(self, text="SESSION RESULTS", bg=PANEL_BG, fg=ACCENT,
                 font=("Segoe UI", 26, "bold")).grid(row=0, column=0, pady=(36, 4))
        _mode_labels = {"tracking_only": "Tracking Only",
                        "codes_only": "Codes Only",
                        "dual_task": "Dual Task"}
        _mode_str = _mode_labels.get(r.get("task_mode", ""), r.get("task_mode", ""))
        tk.Label(self,
                 text=f"{id_header(r['subject_id'], r['session_id'])}   |   Mode: {_mode_str}",
                 bg=PANEL_BG, fg=FG, font=("Segoe UI", 11)).grid(row=1, column=0, pady=(0, 20))

        separator(self).grid(row=2, column=0, sticky="ew", padx=60, pady=(0, 18))

        # Two-column score area
        scores = tk.Frame(self, bg=PANEL_BG)
        scores.grid(row=3, column=0, padx=60, pady=(0, 10))

        def score_block(parent, col, title, items):
            tk.Label(parent, text=title, bg=PANEL_BG, fg=ACCENT2,
                     font=("Segoe UI", 13, "bold")).grid(
                row=0, column=col, padx=30, pady=(0, 10))
            for i, (label, value) in enumerate(items):
                tk.Label(parent, text=label + ":", bg=PANEL_BG, fg="#888899",
                         font=("Segoe UI", 10)).grid(
                    row=i + 1, column=col, sticky="w", padx=30, pady=2)
                tk.Label(parent, text=str(value), bg=PANEL_BG, fg=FG,
                         font=("Segoe UI", 12, "bold")).grid(
                    row=i + 1, column=col + 1, sticky="w", padx=(0, 30), pady=2)

        score_block(scores, 0, "TRACKING",
                    [("Time on Target", f"{r['time_on_target']}%"),
                     ("Avg Error (px)", r["avg_tracking_error"]),
                     ("Max Error (px)", r["max_tracking_error"]),
                     ("Duration (s)", r["duration"])])

        # Visual divider
        tk.Frame(scores, bg=SEP, width=1).grid(
            row=0, column=2, rowspan=5, sticky="ns", padx=10)

        score_block(scores, 3, "SECONDARY TASK",
                    [("Total Codes", r["total_codes"]),
                     ("Correct", r["correct_codes"]),
                     ("Incorrect", r["incorrect_codes"]),
                     ("Accuracy", f"{r['code_accuracy']}%"),
                     ("Avg RT (correct)", f"{r['avg_correct_rt']} s")])

        separator(self).grid(row=4, column=0, sticky="ew", padx=60, pady=18)

        # Save path info
        if "save_dir" in r:
            tk.Label(self,
                     text=f"Data saved to: .../{r['base_name']}",
                     bg=PANEL_BG, fg="#555575",
                     font=("Segoe UI", 9, "italic")).grid(
                row=5, column=0, pady=(0, 18))

        # Buttons
        btn_frame = tk.Frame(self, bg=PANEL_BG)
        btn_frame.grid(row=6, column=0, pady=(0, 40))

        styled_button(btn_frame, "Run Again", self._run_again,
                      width=12, size=11, color=ACCENT2).pack(side="left", padx=8)
        styled_button(btn_frame, "Change Parameters",
                      self.master.show_params,
                      width=20, size=11, color=ACCENT2).pack(side="left", padx=8)
        styled_button(btn_frame, "New Subject",
                      self.master.show_login,
                      width=14, size=11, color=ACCENT2).pack(side="left", padx=8)
        styled_button(btn_frame, "Close",
                      self.master.destroy, width=10, size=11,
                      color="#ff6666").pack(side="left", padx=8)

        s = self.master.scale
        self.master.geometry(f"{int(780*s)}x{int(560*s)}")
        self._centre_window()

    def _centre_window(self):
        s = self.master.scale
        w, h = int(780 * s), int(560 * s)
        sw = self.master.winfo_screenwidth()
        sh = self.master.winfo_screenheight()
        self.master.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def _run_again(self):
        # Advance mode through the standard training sequence
        _next = {"tracking_only": "codes_only",
                 "codes_only": "dual_task",
                 "dual_task": "dual_task"}
        self.master.task_mode.set(_next[self.master.task_mode.get()])
        self.master.show_params()


# ─────────────────────────────────────────────────────────────
#  Entry point
# ─────────────────────────────────────────────────────────────
if __name__ == "__main__":
    app = DualTaskApp()
    app.mainloop()
