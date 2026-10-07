from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin   = Inches(1.25)
section.right_margin  = Inches(1.25)

# ── Helpers ───────────────────────────────────────────────────
def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x66, 0x33)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(2)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x44, 0x88)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x55)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    return p

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    if p.runs:
        p.runs[0].font.size = Pt(11)
    return p

def add_bullet(doc, prefix, rest):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    if prefix:
        r1 = p.add_run(prefix)
        r1.bold = True
        r1.font.size = Pt(11)
    r2 = p.add_run(rest)
    r2.font.size = Pt(11)
    return p

def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run('NOTE: ')
    r1.bold = True
    r1.font.color.rgb = RGBColor(0xAA, 0x55, 0x00)
    r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    r2.font.italic = True
    return p

def add_screenshot(doc, label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('[  SCREENSHOT: ' + label + '  ]')
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)
    return p

def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(2)

def fill_table(tbl, rows_data, header_row=True):
    for i, row in enumerate(rows_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            runs = cell.paragraphs[0].runs
            if runs:
                runs[0].font.size = Pt(10)
                if i == 0 and header_row:
                    runs[0].bold = True

# ══════════════════════════════════════════════════════════════
# TITLE PAGE
# ══════════════════════════════════════════════════════════════
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = t.add_run('STANDARD OPERATING PROCEDURE')
tr.font.size = Pt(22)
tr.font.bold = True
tr.font.color.rgb = RGBColor(0x00, 0x66, 0x33)
t.paragraph_format.space_before = Pt(36)
t.paragraph_format.space_after  = Pt(8)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.add_run('DualTask Application  (v2.1)').font.size = Pt(16)
s.runs[0].font.bold = True
s.paragraph_format.space_after = Pt(4)

s2 = doc.add_paragraph()
s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
s2.add_run('Mouse Tracking & Secondary Coding Task').font.size = Pt(13)
s2.runs[0].font.italic = True
s2.paragraph_format.space_after = Pt(36)

meta = doc.add_table(rows=5, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = 'Table Grid'
meta_data = [
    ('Document Version', 'v2.1'),
    ('Effective Date',   '___________'),
    ('Prepared By',      '___________'),
    ('Reviewed By',      '___________'),
    ('Classification',   'Research Use Only'),
]
for i, (k, v) in enumerate(meta_data):
    meta.cell(i, 0).text = k
    meta.cell(i, 1).text = v
    meta.cell(i, 0).paragraphs[0].runs[0].bold = True
    meta.cell(i, 0).paragraphs[0].runs[0].font.size = Pt(11)
    meta.cell(i, 1).paragraphs[0].runs[0].font.size = Pt(11)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 1. PURPOSE
# ══════════════════════════════════════════════════════════════
add_h1(doc, '1. Purpose')
add_hr(doc)
add_body(doc, (
    'This SOP describes the procedures for operating the DualTask application (v2.1), a '
    'Windows-based psychomotor research tool that measures (1) continuous mouse-tracking '
    'performance, (2) a timed secondary digit-code entry task, or (3) both simultaneously. '
    'The application supports a structured three-trial training sequence and saves all data '
    'to a single Excel workbook per trial for subsequent analysis.'
))

# ══════════════════════════════════════════════════════════════
# 2. SCOPE
# ══════════════════════════════════════════════════════════════
add_h1(doc, '2. Scope')
add_hr(doc)
add_body(doc, (
    'This procedure applies to all research personnel administering the DualTask application to '
    'human subjects. It covers application launch, session setup, task mode selection, task '
    'administration, data verification, and post-session procedures.'
))

# ══════════════════════════════════════════════════════════════
# 3. EQUIPMENT & SYSTEM REQUIREMENTS
# ══════════════════════════════════════════════════════════════
add_h1(doc, '3. Equipment & System Requirements')
add_hr(doc)
add_bullet(doc, '', 'Windows 10 or Windows 11 computer (64-bit)')
add_bullet(doc, '', 'Mouse or trackpad (mouse strongly recommended for tracking task)')
add_bullet(doc, '', 'Full-HD or higher resolution display (application auto-scales to screen resolution)')
add_bullet(doc, '', 'Keyboard (numeric keys used for code entry during task)')
add_bullet(doc, '', 'DualTask_v2.1 application folder containing DualTask_v2.1.exe and _internal\\')
add_bullet(doc, '', 'Designated writable data save folder (default: data\\ subfolder alongside the exe)')
add_note(doc, (
    'The exe must be run from a writable directory (not C:\\Program Files). '
    'Place the DualTask_v2.1 folder in Downloads, Desktop, or a study-specific folder.'
))

# ══════════════════════════════════════════════════════════════
# 4. APPLICATION OVERVIEW
# ══════════════════════════════════════════════════════════════
add_h1(doc, '4. Application Overview')
add_hr(doc)
add_body(doc, 'The application guides the operator and subject through the following sequential screens:')
doc.add_paragraph()

flow_tbl = doc.add_table(rows=7, cols=2)
flow_tbl.style = 'Table Grid'
fill_table(flow_tbl, [
    ('Screen', 'Description'),
    ('1. Login', 'Operator enters Subject ID and Session ID.'),
    ('2. Subjective Ratings',
     'Subject rates sleepiness (Stanford Sleepiness Scale 1–7) and fatigue '
     '(Subjective Fatigue Scale 1–7). Collected once per subject login.'),
    ('3. Task Parameters',
     'Operator reviews settings including Task Mode. All controls locked by default; '
     'passcode required to modify. Task Mode auto-advances through the standard sequence.'),
    ('4. Task (Fullscreen)',
     '5-second countdown followed by the active task(s) for the configured duration.'),
    ('5. Results',
     'Summary performance metrics displayed. Data automatically saved to an Excel workbook.'),
    ('6. Navigation',
     'Operator selects: Run Again (advances task mode) / Change Parameters / New Subject / Close.'),
])
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
# 5. TASK MODES
# ══════════════════════════════════════════════════════════════
add_h1(doc, '5. Task Modes')
add_hr(doc)
add_body(doc, (
    'The application provides three task modes selectable from the Task Parameters screen (requires unlock). '
    'On initial launch the mode defaults to Tracking Only. Pressing Run Again automatically '
    'advances to the next mode in the standard sequence. The sequence resets when a new subject logs in.'
))
doc.add_paragraph()

mode_tbl = doc.add_table(rows=4, cols=3)
mode_tbl.style = 'Table Grid'
fill_table(mode_tbl, [
    ('Mode', 'Active Components', 'Typical Use'),
    ('Tracking Only',
     'Bubble tracking only. No codes presented.',
     'Trial 1 — familiarisation with the tracking task.'),
    ('Codes Only',
     'Code entry only. No bubble or crosshair shown.',
     'Trial 2 — familiarisation with the code entry task in isolation.'),
    ('Dual Task',
     'Both tracking and code entry simultaneously.',
     'Trial 3+ — the primary experimental condition.'),
])
doc.add_paragraph()

add_note(doc, (
    'The Task Mode selector is locked by default. The operator must unlock parameters to change '
    'it manually. Run Again advances the mode automatically without requiring the passcode.'
))

# ══════════════════════════════════════════════════════════════
# 6. STEP-BY-STEP PROCEDURES
# ══════════════════════════════════════════════════════════════
add_h1(doc, '6. Step-by-Step Procedures')
add_hr(doc)

# 6.1
add_h2(doc, '6.1  Launching the Application')
add_body(doc, 'Navigate to the DualTask_v2.1 folder and double-click DualTask_v2.1.exe.')
add_note(doc, (
    'On first launch Windows SmartScreen may show a warning. '
    'Click "More info" then "Run anyway". This is expected for unsigned research software.'
))
add_body(doc, 'The application opens with the Login screen centred on the display.')
add_screenshot(doc, 'Login Screen')

# 6.2
add_h2(doc, '6.2  Screen 1 — Login')
add_body(doc, (
    'The Login screen collects the subject and session identifiers used to name the output files.'
))
add_bullet(doc, '', 'Enter the assigned Subject ID (e.g., 1, 2, P01).')
add_bullet(doc, '', 'Enter the Session ID (e.g., 1, 2, baseline, post).')
add_bullet(doc, '', 'Press Enter or click CONTINUE.')
add_note(doc, (
    'Both fields are required. These values are embedded in the output filenames and cannot '
    'be changed after continuing without logging in as a new subject.'
))

# 6.3
add_h2(doc, '6.3  Screen 2 — Subjective Ratings')
add_body(doc, (
    'The Subjective Ratings screen is shown once per subject login. The subject self-reports '
    'their current alertness and fatigue, and reports their dominant mouse hand, before '
    'beginning any task trials.'
))
add_screenshot(doc, 'Subjective Ratings Screen')

add_h3(doc, 'Stanford Sleepiness Scale (SSS)  —  Hoddes et al., Psychophysiol, 1973')
add_body(doc, 'Ask the subject: "Please select the number that best describes how you feel right now."')
sss_tbl = doc.add_table(rows=8, cols=2)
sss_tbl.style = 'Table Grid'
fill_table(sss_tbl, [
    ('Rating', 'Description'),
    ('1', 'Feeling active and vital; alert; wide awake'),
    ('2', 'Functioning at a high level, but not at peak; able to concentrate'),
    ('3', 'Relaxed; awake; not at full alertness; responsive'),
    ('4', 'A little foggy; not at peak; let down'),
    ('5', 'Fogginess; beginning to lose interest in remaining awake; slowed down'),
    ('6', 'Sleepiness; prefer to be lying down; fighting sleep; woozy'),
    ('7', 'Almost in reverie; sleep onset soon; lost struggle to remain awake'),
])
doc.add_paragraph()

add_h3(doc, 'Subjective Fatigue Scale (SFS)  —  Samm-Perelli Crew Status Check')
add_body(doc, 'Ask the subject: "Now please select the number that best describes your fatigue level."')
sfs_tbl = doc.add_table(rows=8, cols=2)
sfs_tbl.style = 'Table Grid'
fill_table(sfs_tbl, [
    ('Rating', 'Description'),
    ('1', 'Fully alert; Wide awake'),
    ('2', 'Very lively; Responsive; But not at peak'),
    ('3', 'Okay; Somewhat fresh'),
    ('4', 'A little tired; less than peak'),
    ('5', 'Moderately tired; let down'),
    ('6', 'Extremely tired; Very difficult to concentrate'),
    ('7', 'Completely exhausted; Unable to function effectively; Ready to drop'),
])
doc.add_paragraph()

add_h3(doc, 'Dominant Mouse Hand')
add_body(doc, (
    'Ask the subject: "Which hand do you normally use to operate a computer mouse?" '
    'The subject selects Right or Left. This selection is used for two purposes: it is saved '
    'to the output data (mouse_hand, on the Summary tab — see Section 7), and it drives the '
    'wording of the on-screen task instructions on the following Task Parameters screen, which '
    'explicitly tell the subject which hand to use for tracking and which hand to use for code '
    'entry, based on this response.'
))

add_body(doc, (
    'All three selections (Sleepiness, Fatigue, and Mouse Hand) must be made before clicking '
    'CONTINUE. '
    'If the subject runs multiple trials (Run Again / Change Parameters), '
    'the ratings screen will NOT appear again. The original scores and mouse hand are saved '
    'with every trial.'
))

# 6.4
add_h2(doc, '6.4  Screen 3 — Task Parameters')
add_body(doc, (
    'The Task Parameters screen displays the current task configuration. '
    'All controls are locked by default to prevent subjects from modifying them.'
))
add_screenshot(doc, 'Task Parameters Screen')

add_h3(doc, 'Default Configuration')
param_tbl = doc.add_table(rows=7, cols=3)
param_tbl.style = 'Table Grid'
fill_table(param_tbl, [
    ('Parameter', 'Default', 'Description'),
    ('Duration (s)', '100',
     'Total task duration in seconds. With defaults, 18 codes fit within the available window.'),
    ('Number of Codes', '18',
     'Total 4-digit codes to present. Codes appear at fixed 5-second intervals after a 5-second '
     'initial buffer. Excess codes beyond the task duration are silently ignored.'),
    ('Code Interval +/- (s)', '5',
     'Display window per code (seconds). Each code disappears after 5 s; marked MISSED if unanswered.'),
    ('Max Speed (px/tick)', '6.0', 'Maximum bubble movement speed.'),
    ('Min Speed (px/tick)', '2.0', 'Minimum bubble movement speed.'),
    ('Mouse Gain', '1.0',
     'Cursor sensitivity (1.0 = 1:1 mapping; <1.0 = reduced range; >1.0 = amplified).'),
])
doc.add_paragraph()

add_h3(doc, 'Task Mode')
add_body(doc, (
    'The Task Mode radio buttons determine which components are active for the upcoming trial. '
    'The mode defaults to Tracking Only on first launch and advances automatically when '
    'Run Again is pressed. See Section 5 for full mode descriptions.'
))
doc.add_paragraph()

add_h3(doc, 'Unlocking Parameters (Operator Only)')
add_bullet(doc, '', 'Click the Lock / Unlock Params button.')
add_bullet(doc, '', 'Enter the operator passcode when prompted and press Enter or click Unlock.')
add_bullet(doc, '', 'All fields and the Task Mode selector become editable.')
add_bullet(doc, '', 'Make required changes, then click START TASK.')
add_note(doc, (
    'Do not share the passcode with subjects. '
    'The lock re-engages automatically when navigating to Change Parameters or New Subject.'
))

add_h3(doc, 'Save Folder')
add_body(doc, (
    'The Save Folder field shows where output files will be written. '
    'Click Browse to select a different location. '
    'Ensure the folder exists and is writable before starting the task.'
))
add_body(doc, 'Click START TASK to proceed.')

# 6.5
add_h2(doc, '6.5  Screen 4 — Task (Fullscreen)')
add_body(doc, (
    'The application switches to fullscreen with a black background. '
    'A 5-second countdown is displayed before the task begins. '
    'The cursor is repositioned to the bubble location at task start.'
))
add_screenshot(doc, 'Countdown Screen')
add_screenshot(doc, 'Task Screen (active trial)')

add_note(doc, (
    'Starting in v2.0, the on-screen instructions on the Task Parameters screen (Section 6.4) '
    'are generated dynamically from the subject’s Dominant Mouse Hand response (Section 6.3): '
    'the tracking instructions name the subject’s selected mouse hand explicitly (e.g., '
    '"use your Right hand"), and the code-entry instructions name the opposite hand. The '
    'descriptions below show the underlying instruction text with generic hand references.'
))

add_h3(doc, 'Subject Instructions — Tracking Only (Trial 1)')
p1 = doc.add_paragraph(style='List Number')
p1.add_run('Tracking Task: ').bold = True
p1.runs[0].font.size = Pt(11)
p1.add_run(
    'A green circle (bubble) will move around the screen. '
    'Use the mouse to keep the white crosshair cursor inside the bubble at all times. '
    'Try to stay inside as much as possible for the full duration. Use the hand you reported '
    'as your mouse hand on the Subjective Ratings screen.'
).font.size = Pt(11)

add_h3(doc, 'Subject Instructions — Codes Only (Trial 2)')
p2 = doc.add_paragraph(style='List Number')
p2.add_run('Code Entry Task: ').bold = True
p2.runs[0].font.size = Pt(11)
p2.add_run(
    'A 4-digit code will appear at the top of the screen (e.g., Enter: 4829). '
    'Type the four digits as quickly and accurately as possible using the hand opposite your '
    'reported mouse hand. '
    'The code disappears automatically after 5 seconds whether or not you have responded. '
    'A new code will appear every 5 seconds.'
).font.size = Pt(11)

add_h3(doc, 'Subject Instructions — Dual Task (Trial 3)')
p3a = doc.add_paragraph(style='List Number')
p3a.add_run('Tracking Task: ').bold = True
p3a.runs[0].font.size = Pt(11)
p3a.add_run(
    'Keep the crosshair inside the moving bubble using the same hand as the Tracking Only task.'
).font.size = Pt(11)

p3b = doc.add_paragraph(style='List Number')
p3b.add_run('Code Entry Task: ').bold = True
p3b.runs[0].font.size = Pt(11)
p3b.add_run(
    'When a code appears at the top of the screen, type the four digits quickly using the same '
    'hand as the Codes Only task. '
    'The code disappears after 5 seconds; continue tracking immediately after entering.'
).font.size = Pt(11)

p3c = doc.add_paragraph(style='List Number')
p3c.add_run('Priority: ').bold = True
p3c.runs[0].font.size = Pt(11)
p3c.add_run(
    'Give equal priority to both tasks. Do not sacrifice tracking to enter a code, and vice versa.'
).font.size = Pt(11)
doc.add_paragraph()

add_h3(doc, 'During the Task — Operator Reference')
add_bullet(doc, '', 'Thin blue progress bar fills left-to-right at the top of the screen.')
add_bullet(doc, '', 'Remaining time (seconds) shown in the top-right corner.')
add_bullet(doc, 'Code prompt: ',
           'Display shows:  Enter: XXXX      Your input: ____')
add_bullet(doc, '', 'Digits typed replace underscores. Code auto-submits on the 4th digit.')
add_bullet(doc, '', 'Each code is visible for exactly 5 seconds. Unanswered codes are recorded as MISSED.')
add_bullet(doc, '', 'Press Escape to abort the task early if needed.')

# 6.6
add_h2(doc, '6.6  Screen 5 — Results')
add_body(doc, (
    'After the task ends the Results screen displays a performance summary. '
    'Data is automatically saved to the configured Save Folder as an Excel workbook.'
))
add_screenshot(doc, 'Results Screen')

add_h3(doc, 'Tracking Metrics  (Tracking Only and Dual Task modes)')
add_bullet(doc, 'Time on Target (%): ',
           'Percentage of task time the cursor was within the bubble.')
add_bullet(doc, 'Avg Error (px): ',
           'Mean pixel distance between cursor and bubble centre (~50 Hz samples).')
add_bullet(doc, 'Max Error (px): ', 'Peak pixel distance recorded during the trial.')

add_h3(doc, 'Secondary Task Metrics  (Codes Only and Dual Task modes)')
add_bullet(doc, 'Total Codes: ', 'Number of codes presented during the trial.')
add_bullet(doc, 'Correct: ', 'Codes entered exactly correctly within the 5-second window.')
add_bullet(doc, 'Incorrect: ',
           'Codes answered incorrectly or only partially entered before the window expired.')
add_bullet(doc, 'Accuracy (%): ', 'Percentage of total codes answered correctly.')
add_bullet(doc, 'Avg RT (correct): ',
           'Mean reaction time in seconds from code onset to 4th digit, for correct responses only.')

add_h3(doc, 'Navigation Options After Results')
nav_tbl = doc.add_table(rows=5, cols=2)
nav_tbl.style = 'Table Grid'
fill_table(nav_tbl, [
    ('Button', 'Action'),
    ('Run Again',
     'Start the next trial. Task Mode auto-advances: Tracking Only → Codes Only → Dual Task. '
     'Subjective ratings NOT repeated.'),
    ('Change Parameters',
     'Return to Parameters screen. Task Mode unchanged. Subjective ratings NOT repeated.'),
    ('New Subject',
     'Return to Login. Task Mode resets to Tracking Only. Subjective ratings WILL be shown.'),
    ('Close', 'Exit the application.'),
])
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
# 7. OUTPUT FILES
# ══════════════════════════════════════════════════════════════
add_h1(doc, '7. Output Files')
add_hr(doc)
add_body(doc, (
    'One Excel workbook (.xlsx) is created automatically at the end of each trial, named:'
))
p = doc.add_paragraph()
r = p.add_run(
    '    {SubjectID}_session{SessionID}_{task_mode}_trial{N}.xlsx'
)
r.font.name = 'Courier New'
r.font.size = Pt(10)
doc.add_paragraph()

add_body(doc, 'Examples:')
p2 = doc.add_paragraph()
r2 = p2.add_run(
    '    P01_session1_tracking_only_trial1.xlsx\n'
    '    P01_session1_codes_only_trial2.xlsx\n'
    '    P01_session1_dual_task_trial3.xlsx'
)
r2.font.name = 'Courier New'
r2.font.size = Pt(10)
doc.add_paragraph()

add_note(doc, (
    'The trial number increments each time a trial completes and resets to 1 when a new subject '
    'logs in. If a session is repeated (e.g., re-running from Change Parameters), '
    'the trial number continues from where it left off.'
))

add_h3(doc, 'Data Tab  (event-log format, one row per event)')
add_body(doc, (
    'Contains all time-stamped events for the trial. For Dual Task mode, tracking and code events '
    'are interleaved in chronological order. Researchers can filter by event_type to isolate '
    'each stream.'
))
data_col_tbl = doc.add_table(rows=12, cols=2)
data_col_tbl.style = 'Table Grid'
fill_table(data_col_tbl, [
    ('Column', 'Description'),
    ('elapsed_s',
     'Time of the event since task start (seconds).'),
    ('event_type',
     '"tracking" for motion samples; "code" for code presentation events.'),
    ('cursor_x, cursor_y',
     'Crosshair position in screen pixels. Filled for tracking rows; blank for code rows.'),
    ('bubble_x, bubble_y',
     'Bubble centre position in screen pixels. Filled for tracking rows; blank for code rows.'),
    ('error_px',
     'Euclidean distance between cursor and bubble centre (pixels). Tracking rows only.'),
    ('code',
     'The 4-digit code presented. Code rows only.'),
    ('response',
     'Digits entered by the subject. "MISSED" if the 5-second window expired with no entry; '
     '"NO_RESPONSE" if the task ended while the code was still active. Code rows only.'),
    ('rt_s',
     'Reaction time from code onset to 4th digit entry (seconds). Code rows only.'),
    ('correct',
     '1 if the response matched the code exactly, 0 otherwise. Code rows only.'),
])
doc.add_paragraph()

add_h3(doc, 'Summary Tab  (key-value format)')
add_body(doc, (
    'Upper section — session results: subject_id, session_id, task_mode, sss_scale, sfs_scale, '
    'mouse_hand, duration, avg_tracking_error, max_tracking_error, time_on_target, total_codes, '
    'correct_codes, incorrect_codes, code_accuracy, avg_correct_rt.'
))
add_body(doc, (
    'Lower section — task parameters used for the trial: duration, num_alarms, alarm_variance, '
    'max_speed, min_speed, mouse_gain, task_mode, buzzer_on, bubble_radius_px, tail_length, '
    'tick_ms, code_digits, countdown_from.'
))

# ══════════════════════════════════════════════════════════════
# 8. TROUBLESHOOTING
# ══════════════════════════════════════════════════════════════
add_h1(doc, '8. Troubleshooting')
add_hr(doc)
ts_tbl = doc.add_table(rows=9, cols=2)
ts_tbl.style = 'Table Grid'
fill_table(ts_tbl, [
    ('Issue', 'Resolution'),
    ('Application blocked by antivirus',
     'Add the DualTask_v2.1 folder as an antivirus exception.'),
    ('Taskbar icon shows feather instead of logo',
     'Unpin and re-pin the application to the taskbar after replacing the exe.'),
    ('Mouse cursor restricted to part of screen',
     'Ensure the application is not running via remote desktop. Restart the application.'),
    ('No buzzer sound on code appearance',
     'Confirm Play buzzer on code is checked in Task Parameters (requires unlock). '
     'Buzzer is off by default.'),
    ('Output file not saved / error on close',
     'Confirm the Save Folder is writable and not in C:\\Program Files. '
     'Ensure openpyxl is bundled correctly in the _internal folder.'),
    ('Application appears small on high-DPI display',
     'Confirm Windows Display Scaling is set correctly. The application auto-scales to physical '
     'screen resolution.'),
    ('Subjective ratings text cut off',
     'The Subjective Ratings window dynamically resizes to fit content. If this persists, '
     'increase the Windows display font size and relaunch.'),
    ('Task mode not advancing on Run Again',
     'Mode only advances automatically via the Run Again button. '
     'Change Parameters and New Subject do not advance the mode.'),
])
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
# 9. DATA HANDLING
# ══════════════════════════════════════════════════════════════
add_h1(doc, '9. Data Handling & Storage')
add_hr(doc)
add_bullet(doc, '', 'Back up .xlsx files to a secure server or cloud drive immediately after each session.')
add_bullet(doc, '', (
    'File names encode subject ID, session, task mode, and trial number — do not rename output files.'))
add_bullet(doc, '', 'Retain raw workbooks in their original form and perform all analyses on copies.')
add_bullet(doc, '', "Follow your institution's data management plan for subject privacy and retention schedules.")

# ══════════════════════════════════════════════════════════════
# 10. VERSION HISTORY
# ══════════════════════════════════════════════════════════════
add_h1(doc, '10. Version History')
add_hr(doc)
vh_tbl = doc.add_table(rows=13, cols=3)
vh_tbl.style = 'Table Grid'
fill_table(vh_tbl, [
    ('Version', 'Date', 'Changes'),
    ('v1.0', 'Apr 2026',
     'Initial release. Login, Params, Task, Results screens. Tracking + code entry. CSV output.'),
    ('v1.1', 'Apr 2026',
     'Added NAS logo icon. Fixed DLL loading issue on distribution machines.'),
    ('v1.2', 'Apr 2026',
     'Auto-scaling for high-DPI displays (Surface Pro compatibility).'),
    ('v1.3', 'Apr 2026',
     'Countdown overlap fix. Parameter passcode lockout added.'),
    ('v1.4', 'Apr 2026',
     'Countdown elements repositioned by screen fraction to prevent overlap at any resolution.'),
    ('v1.5', 'May 2026',
     'Subjective Ratings screen (SSS + SFS) added. Code suppression in final 3 seconds of task.'),
    ('v1.6', 'Jun 2026',
     'Renamed Missed Codes to Incorrect Codes. Reserved top task area so cursor and bubble do not '
     'overlap code prompt. Increased Subjective Ratings screen text size.'),
    ('v1.7', 'Jun 2026',
     'Mouse cursor starts on target during countdown. Summary CSV labels subjective scales as '
     'sss_scale/sfs_scale and appends task parameters.'),
    ('v1.8', 'Jun 2026',
     'Three task modes: Tracking Only (default), Codes Only, Dual Task — auto-advance sequence on '
     'Run Again. Fixed 5-second code intervals (5 s initial buffer, 5 s window per code; MISSED if '
     'unanswered). Buzzer off by default. Subjective ratings window dynamically resizes to prevent '
     'text clipping. Output changed from separate CSV files to single Excel workbook per trial '
     '(Data tab + Summary tab) named by task mode and trial number.'),
    ('v1.9', 'Jun 2026',
     'Internal build/packaging update.'),
    ('v2.0', 'Jun 2026',
     'Subjective Ratings screen now also asks the subject which hand they normally use to '
     'operate a computer mouse (Right/Left), recorded as mouse_hand on the Summary tab. '
     'Tracking, Codes Only, and Dual Task instructions now reference the subject’s selected '
     'hand explicitly (e.g., "use your Right hand") instead of generic phrasing. Build '
     'packaging fixed to stop bundling unrelated third-party libraries pulled in accidentally '
     'by the previous PyInstaller spec, reducing the installed footprint from 300+ MB to ~30 MB.'),
    ('v2.1', 'Jul 2026',
     'Fixed the Right/Left mouse-hand buttons on the Subjective Ratings screen being displayed '
     'in swapped positions (Right rendered on the left, Left on the right).'),
])
doc.add_paragraph()

# ── Save ──────────────────────────────────────────────────────
out = r'c:\Users\ambolli1\Documents\Python\DualTask_AMB_build\DualTask_v2.1_SOP.docx'
doc.save(out)
print('Saved: ' + out)
